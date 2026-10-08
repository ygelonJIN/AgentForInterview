"""把用户上传的常见文档转换成可写入 RAG 的文本记录。"""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional

SUPPORTED_DOCUMENT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".jsonl",
    ".pdf",
    ".docx",
    ".xlsx",
    ".xls",
    ".csv",
}
SUPPORTED_IMPORT_SUFFIXES = SUPPORTED_DOCUMENT_SUFFIXES


def _document(page_content: str, metadata: Optional[Mapping[str, Any]] = None):
    from langchain_core.documents import Document

    return Document(page_content=str(page_content), metadata=dict(metadata or {}))


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def _load_jsonl(path: Path, shared: Mapping[str, Any]) -> List[Any]:
    records: List[Any] = []
    for line_number, line in enumerate(_read_text(path).splitlines(), start=1):
        if not line.strip():
            continue
        try:
            payload = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ValueError(f"{path}:{line_number} JSONL 解析失败: {exc}") from exc
        if not isinstance(payload, dict):
            raise ValueError(f"{path}:{line_number} 每行必须是 JSON 对象")
        content = payload.get("content", payload.get("text", ""))
        if not isinstance(content, str) or not content.strip():
            raise ValueError(f"{path}:{line_number} 缺少非空 content/text")
        metadata = dict(payload.get("metadata") or {})
        metadata.update({
            key: value for key, value in payload.items()
            if key not in {"content", "text", "metadata"} and value is not None
        })
        metadata.update(shared)
        records.append(_document(content, metadata))
    return records


def _load_csv(path: Path, shared: Mapping[str, Any]) -> List[Any]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError(f"{path} 缺少表头")
        records = []
        for row_number, row in enumerate(reader, start=2):
            values = {key: value for key, value in row.items() if key and value not in (None, "")}
            content = "\n".join(f"{key}: {value}" for key, value in values.items())
            if not content.strip():
                continue
            metadata = dict(shared)
            metadata.update({"row_number": row_number})
            metadata.update({f"csv_{key}": value for key, value in values.items()})
            records.append(_document(content, metadata))
        return records


def _load_pdf(path: Path, shared: Mapping[str, Any]) -> List[Any]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValueError("导入 PDF 需要安装 pypdf，请执行 pip install -r requirements.txt") from exc
    reader = PdfReader(str(path))
    records = []
    for page_number, page in enumerate(reader.pages, start=1):
        content = (page.extract_text() or "").strip()
        if content:
            metadata = dict(shared)
            metadata.update({"page_number": page_number})
            records.append(_document(content, metadata))
    if not records:
        raise ValueError(f"{path} 没有可提取的文本")
    return records


def _load_docx(path: Path, shared: Mapping[str, Any]) -> List[Any]:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise ValueError("导入 DOCX 需要安装 python-docx，请执行 pip install -r requirements.txt") from exc
    doc = DocxDocument(str(path))
    blocks: List[str] = []
    blocks.extend(paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip())
    for table_index, table in enumerate(doc.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append(" | ".join(values))
    content = "\n".join(blocks).strip()
    if not content:
        raise ValueError(f"{path} 没有可提取的文本")
    metadata = dict(shared)
    metadata.update({"document_type": "docx"})
    return [_document(content, metadata)]


def _load_excel(path: Path, shared: Mapping[str, Any]) -> List[Any]:
    try:
        import pandas as pd
    except ImportError as exc:
        raise ValueError("导入 Excel 需要安装 pandas/openpyxl，请执行 pip install -r requirements.txt") from exc
    try:
        sheets = pd.read_excel(path, sheet_name=None, dtype=object)
    except Exception as exc:
        raise ValueError(f"{path} Excel 解析失败: {exc}") from exc
    records: List[Any] = []
    for sheet_name, frame in sheets.items():
        for row_number, row in enumerate(frame.to_dict(orient="records"), start=2):
            values = {
                str(key): ("" if value is None or (isinstance(value, float) and value != value) else str(value))
                for key, value in row.items()
                if value is not None and str(value).strip()
            }
            content = "\n".join(f"{key}: {value}" for key, value in values.items())
            if not content.strip():
                continue
            metadata = dict(shared)
            metadata.update({"sheet": str(sheet_name), "row_number": row_number})
            records.append(_document(content, metadata))
    if not records:
        raise ValueError(f"{path} 没有可提取的文本")
    return records


def load_document_records(
    file_path: str | Path,
    source_type: str = "guides",
    extra_metadata: Optional[Mapping[str, Any]] = None,
) -> List[Any]:
    """加载单个文件为 Document 列表，所有格式共享来源元数据。"""
    path = Path(file_path).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(str(path))
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_DOCUMENT_SUFFIXES:
        raise ValueError(f"不支持的文档格式: {suffix or '(无扩展名)'}")
    shared: Dict[str, Any] = {
        "source": str(path),
        "source_type": source_type,
        "owner": "public",
        **dict(extra_metadata or {}),
    }
    if suffix == ".jsonl":
        records = _load_jsonl(path, shared)
    elif suffix == ".csv":
        records = _load_csv(path, shared)
    elif suffix == ".pdf":
        records = _load_pdf(path, shared)
    elif suffix == ".docx":
        records = _load_docx(path, shared)
    elif suffix in {".xlsx", ".xls"}:
        records = _load_excel(path, shared)
    else:
        records = [_document(_read_text(path), shared)]
    if not records:
        raise ValueError(f"{path} 没有可导入的内容")
    return records


def records_to_dicts(records: Iterable[Any]) -> List[Dict[str, Any]]:
    return [
        {"content": record.page_content, "metadata": dict(record.metadata or {})}
        for record in records
    ]
