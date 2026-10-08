from pathlib import Path

import pytest

from app.document_import import SUPPORTED_DOCUMENT_SUFFIXES, load_document_records


def _write_minimal_pdf(path: Path, text: str) -> None:
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode())
        payload.extend(body)
        payload.extend(b"\nendobj\n")
    xref = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode())
    payload.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(payload)


def test_supported_suffixes_include_office_and_pdf():
    assert {".pdf", ".docx", ".xlsx", ".xls", ".csv"} <= SUPPORTED_DOCUMENT_SUFFIXES


def test_load_pdf_docx_and_excel_records(tmp_path):
    pdf_path = tmp_path / "guide.pdf"
    _write_minimal_pdf(pdf_path, "Hello PDF")
    pdf_records = load_document_records(pdf_path, "guides", {"tenant": "demo"})
    assert "Hello PDF" in pdf_records[0].page_content
    assert pdf_records[0].metadata["page_number"] == 1
    assert pdf_records[0].metadata["tenant"] == "demo"

    from docx import Document as DocxDocument

    docx_path = tmp_path / "guide.docx"
    doc = DocxDocument()
    doc.add_paragraph("DOCX 正文")
    doc.save(docx_path)
    docx_records = load_document_records(docx_path, "guides")
    assert docx_records[0].page_content == "DOCX 正文"

    from openpyxl import Workbook

    xlsx_path = tmp_path / "guide.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "活动"
    sheet.append(["名称", "地点"])
    sheet.append(["西湖徒步", "杭州"])
    workbook.save(xlsx_path)
    xlsx_records = load_document_records(xlsx_path, "activities")
    assert "西湖徒步" in xlsx_records[0].page_content
    assert xlsx_records[0].metadata["sheet"] == "活动"
    assert xlsx_records[0].metadata["row_number"] == 2


def test_unknown_document_format_is_rejected(tmp_path):
    path = tmp_path / "data.bin"
    path.write_bytes(b"binary")
    with pytest.raises(ValueError, match="不支持"):
        load_document_records(path)
